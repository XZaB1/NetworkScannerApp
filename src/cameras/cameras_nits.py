import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext, filedialog
import threading
import json
import os
import ipaddress
import subprocess
import re
from ipaddress import ip_address
import pandas as pd
from utils.path_manager import PathManager
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class NitsWindow:
    def __init__(self, master):
        self.master = master
        self.master.title("Сканирование камер НИЦ")
        self.master.geometry("1000x800")
        
        # Параметры по умолчанию
        self.snmp_community = 'public'
        self.snmp_version = '2c'
        self.stop_scan = False
        self.running = False
        self.ip_ranges = []
        self.selected_ranges = []
        self.data = []
        
        self.create_widgets()
        self.load_ranges()
        self.create_menu()
    
    def create_menu(self):
        menubar = tk.Menu(self.master)
        
        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="Не работает?", command=self.show_instructions)
        file_menu.add_separator()
        menubar.add_cascade(label="Вопросы", menu=file_menu)

        settings_menu = tk.Menu(menubar, tearoff=0)
        settings_menu.add_command(label="Настройки SNMP", command=self.snmp_settings_dialog)
        menubar.add_cascade(label="Настройки", menu=settings_menu)

        self.master.config(menu=menubar)
    
    def show_instructions(self):
        messagebox.showinfo("Не работает?", "Убедитесь, что у вас установлен SNMP")
    
    def snmp_settings_dialog(self):
        dialog = tk.Toplevel(self.master)
        dialog.title("Настройки SNMP")
        dialog.geometry("300x200")
        
        ttk.Label(dialog, text="Community:").pack(pady=5)
        self.community_entry = ttk.Entry(dialog)
        self.community_entry.insert(0, self.snmp_community)
        self.community_entry.pack(pady=5)
        
        ttk.Label(dialog, text="Версия SNMP:").pack(pady=5)
        self.version_combobox = ttk.Combobox(dialog, values=["1", "2c", "3"])
        self.version_combobox.set(self.snmp_version)
        self.version_combobox.pack(pady=5)
        
        def save_settings():
            self.snmp_community = self.community_entry.get()
            self.snmp_version = self.version_combobox.get()
            dialog.destroy()
        
        ttk.Button(dialog, text="Сохранить", command=save_settings).pack(pady=10)
    
    def create_widgets(self):
        """Создание элементов интерфейса"""
        control_frame = ttk.Frame(self.master)
        control_frame.pack(pady=10, fill=tk.X)
        
        range_frame = ttk.LabelFrame(self.master, text="Управление диапазонами IP")
        range_frame.pack(pady=5, padx=10, fill=tk.BOTH, expand=True)
        
        self.canvas = tk.Canvas(range_frame)
        scrollbar = ttk.Scrollbar(range_frame, orient="vertical", command=self.canvas.yview)
        self.scrollable_frame = ttk.Frame(self.canvas)
        
        self.scrollable_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=scrollbar.set)
        
        self.canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # Элементы управления
        self.start_btn = ttk.Button(control_frame, text="Начать сканирование", command=self.start_scan_thread)
        self.start_btn.pack(side=tk.LEFT, padx=5)
        
        self.stop_btn = ttk.Button(control_frame, text="Остановить", command=self.stop_scan_process, state=tk.DISABLED)
        self.stop_btn.pack(side=tk.LEFT, padx=5)
        
        ttk.Button(control_frame, text="Добавить диапазон", command=self.add_range_dialog).pack(side=tk.RIGHT, padx=5)
        ttk.Button(control_frame, text="Удалить выбранное", command=self.remove_selected_ranges).pack(side=tk.RIGHT, padx=5)
        
        # Прогресс-бар
        self.progress = ttk.Progressbar(self.master, orient="horizontal", mode="determinate")
        self.progress.pack(fill=tk.X, padx=10, pady=5)
        
        # Лог и результаты
        notebook = ttk.Notebook(self.master)
        
        log_frame = ttk.Frame(notebook)
        self.log = scrolledtext.ScrolledText(log_frame, height=15, wrap=tk.WORD)
        self.log.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        notebook.add(log_frame, text="Лог выполнения")
        
        result_frame = ttk.Frame(notebook)
        self.result_text = scrolledtext.ScrolledText(result_frame, height=15, wrap=tk.WORD)
        self.result_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        notebook.add(result_frame, text="Результаты")
        
        notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        save_frame = ttk.Frame(self.master)
        save_frame.pack(pady=10)
        
        ttk.Button(save_frame, text="Сохранить в TXT", command=self.save_to_txt).pack(side=tk.LEFT, padx=5)
        ttk.Button(save_frame, text="Сохранить в Excel", command=self.save_to_excel).pack(side=tk.LEFT, padx=5)
        
        self.update_range_checkboxes()
    
    def load_ranges(self):
        """Загрузка диапазонов IP из файла"""
        try:
            os.makedirs("config", exist_ok=True)
            if os.path.exists("config/ranges_nits.json"):
                with open("config/ranges_nits.json", "r") as f:
                    self.ip_ranges = json.load(f)
            else:
                self.ip_ranges = [
                    ["", ""],
                    ["", ""],
                    ["", ""],
                    ["", ""],
                    ["", ""]
                ]
                self.save_ranges()
            
            self.selected_ranges = self.ip_ranges.copy()
            self.update_range_checkboxes()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить диапазоны: {str(e)}")
    
    def save_ranges(self):
        """Сохранение диапазонов в файл"""
        try:
            with open("config/ranges_nits.json", "w") as f:
                json.dump(self.ip_ranges, f, indent=4)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить диапазоны: {str(e)}")
    
    def update_range_checkboxes(self):
        """Обновление чекбоксов диапазонов"""
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        
        self.range_vars = []
        for i, (start, end) in enumerate(self.ip_ranges, 1):
            var = tk.BooleanVar(value=(start, end) in self.selected_ranges)
            self.range_vars.append((var, (start, end)))
            
            chk = ttk.Checkbutton(
                self.scrollable_frame,
                text=f"Диапазон {i}: {start} — {end}",
                variable=var,
                command=lambda v=var, r=(start, end): self.toggle_range(v, r)
            )
            chk.pack(anchor="w", padx=5, pady=2)
    
    def toggle_range(self, var, range):
        """Переключение состояния диапазона"""
        if var.get():
            if range not in self.selected_ranges:
                self.selected_ranges.append(range)
        else:
            if range in self.selected_ranges:
                self.selected_ranges.remove(range)
    
    def add_range_dialog(self):
        """Диалог добавления нового диапазона"""
        dialog = tk.Toplevel(self.master)
        dialog.title("Добавить диапазон IP")
        dialog.geometry("400x200")
        dialog.resizable(False, False)
        
        ttk.Label(dialog, text="Начальный IP:").pack(pady=5)
        start_entry = ttk.Entry(dialog)
        start_entry.pack(pady=5)
        
        ttk.Label(dialog, text="Конечный IP:").pack(pady=5)
        end_entry = ttk.Entry(dialog)
        end_entry.pack(pady=5)
        
        def save_range():
            start = start_entry.get().strip()
            end = end_entry.get().strip()
            
            try:
                ipaddress.ip_address(start)
                ipaddress.ip_address(end)
                
                new_range = [start, end]
                if new_range not in self.ip_ranges:
                    self.ip_ranges.append(new_range)
                    self.selected_ranges.append(new_range)
                    self.save_ranges()
                    self.update_range_checkboxes()
                    dialog.destroy()
                else:
                    messagebox.showwarning("Внимание", "Такой диапазон уже существует")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Некорректные IP-адреса: {str(e)}")
        
        ttk.Button(dialog, text="Добавить", command=save_range).pack(pady=10)
        start_entry.focus()
    
    def remove_selected_ranges(self):
        """Удаление выбранных диапазонов"""
        ranges_to_remove = []
        for var, rng in self.range_vars:
            if var.get():
                ranges_to_remove.append(list(rng))
        
        if not ranges_to_remove:
            messagebox.showwarning("Внимание", "Не выбраны диапазоны для удаления")
            return
        
        if not messagebox.askyesno("Подтверждение", f"Удалить {len(ranges_to_remove)} выбранных диапазонов?"):
            return
        
        self.ip_ranges = [r for r in self.ip_ranges if r not in ranges_to_remove]
        self.selected_ranges = [r for r in self.selected_ranges if r not in ranges_to_remove]
        self.save_ranges()
        self.update_range_checkboxes()
        messagebox.showinfo("Успех", f"Удалено {len(ranges_to_remove)} диапазонов")
    
    def start_scan_thread(self):
        """Запуск сканирования в отдельном потоке"""
        if not self.selected_ranges:
            messagebox.showwarning("Ошибка", "Выберите хотя бы один диапазон!")
            return
        
        if self.running:
            messagebox.showwarning("Внимание", "Сканирование уже выполняется!")
            return
        
        self.stop_scan = False
        self.running = True
        self.data = []
        self.log.delete("1.0", tk.END)
        self.result_text.delete("1.0", tk.END)
        
        self.start_btn.config(state=tk.DISABLED)
        self.stop_btn.config(state=tk.NORMAL)
        
        self.scan_thread = threading.Thread(target=self.run_scan, daemon=True)
        self.scan_thread.start()
    
    def stop_scan_process(self):
        """Остановка сканирования"""
        if not self.running:
            return
        
        self.stop_scan = True
        self.running = False
        self.log.insert(tk.END, "\n[INFO] Сканирование остановлено пользователем\n")
        self.start_btn.config(state=tk.NORMAL)
        self.stop_btn.config(state=tk.DISABLED)
    
    def run_scan(self):
        """Основная функция сканирования"""
        try:
            total_ips = sum(
                int(ipaddress.ip_address(end)) - int(ipaddress.ip_address(start)) + 1
                for start, end in self.selected_ranges
            )
            
            self.progress.config(maximum=total_ips)
            processed = 0
            
            for start, end in self.selected_ranges:
                start_ip = ipaddress.ip_address(start)
                end_ip = ipaddress.ip_address(end)
                
                for ip_int in range(int(start_ip), int(end_ip) + 1):
                    if self.stop_scan or not self.running:
                        self.stop_btn.config(state=tk.DISABLED)
                        self.start_btn.config(state=tk.NORMAL)
                        self.running = False
                        return
                    
                    ip_str = str(ipaddress.ip_address(ip_int))
                    self.log.insert(tk.END, f"[{processed+1}/{total_ips}] Проверка {ip_str}... ")
                    self.log.see(tk.END)
                    self.master.update()
                    
                    camera_info = self.get_camera_info(ip_str)
                    if camera_info:
                        self.data.append(camera_info)
                        self.result_text.insert(tk.END, f"Найдена камера: {ip_str}\n")
                        self.log.insert(tk.END, "OK - камера найдена\n")
                    else:
                        self.log.insert(tk.END, "Нет ответа\n")
                    
                    processed += 1
                    self.progress.config(value=processed)
                    self.master.update()
            
            self.start_btn.config(state=tk.NORMAL)
            self.stop_btn.config(state=tk.DISABLED)
            self.log.insert(tk.END, f"\n[INFO] Сканирование завершено! Найдено камер: {len(self.data)}\n")
            self.running = False
            
        except Exception as e:
            self.log.insert(tk.END, f"\n[ERROR] Критическая ошибка: {str(e)}\n")
            self.start_btn.config(state=tk.NORMAL)
            self.stop_btn.config(state=tk.DISABLED)
            self.running = False
    
    def get_camera_info(self, ip):
        """Получение информации о камере НИЦ через SNMP"""
        try:
            # Указываем полный путь к snmpwalk
            snmpwalk_path = r"C:\usr\bin\snmpwalk.exe"  # Или ваш реальный путь
            
            command = [
                snmpwalk_path,
                "-v", self.snmp_version,
                "-c", self.snmp_community,
                ip,
                "1.3.6.1"
            ]
            
            # Добавляем параметр shell=True для Windows
            result = subprocess.run(
                command,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=5,
                shell=True
            )
            
            if result.returncode != 0:
                error_msg = result.stderr.strip()
                self.log.insert(tk.END, f"Ошибка SNMP: {error_msg}\n")
                return None
            
            data_patterns = {
                "Модель": r'SNMPv2-SMI::enterprises\.31\.1\.1\.0\s*=\s*STRING:\s*"(.*?)"',
                "MAC адрес": r'SNMPv2-SMI::enterprises\.31\.1\.4\.0\s*=\s*STRING:\s*"(.*?)"',
                "Версия ПО": r'SNMPv2-SMI::enterprises\.31\.1\.3\.0\s*=\s*STRING:\s*"(.*?)"',
                "Серийный номер": r'SNMPv2-SMI::enterprises\.31\.1\.5\.0\s*=\s*STRING:\s*"(.*?)"',
                "Аппаратная версия": r'SNMPv2-SMI::enterprises\.31\.1\.2\.0\s*=\s*STRING:\s*"(.*?)"'
            }
            
            extracted_data = {"IP": ip, "Тип устройства": "НИЦ"}
            for key, pattern in data_patterns.items():
                match = re.search(pattern, result.stdout)
                if match:
                    extracted_data[key] = match.group(1)
                else:
                    extracted_data[key] = "Не найден"
            
            # Пропускаем устройства без MAC-адреса
            if extracted_data["MAC адрес"] == "Не найден":
                return None
                
            return extracted_data
            
        except subprocess.TimeoutExpired:
            self.log.insert(tk.END, "Таймаут SNMP запроса\n")
            return None
        except Exception as e:
            self.log.insert(tk.END, f"Ошибка: {str(e)}\n")
            return None
    
    def save_to_txt(self):
        """Сохранение результатов в TXT"""
        if not self.data:
            messagebox.showwarning("Ошибка", "Нет данных для сохранения!")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt")],
            initialfile="nits_scan_results.txt"
        )
        
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as fw:
                    for camera in self.data:
                        fw.write(f"Информация о камере {camera['IP']}:\n")
                        fw.write(f"Модель: {camera['Модель']}\n")
                        fw.write(f"Версия ПО: {camera['Версия ПО']}\n")
                        fw.write(f"Аппаратная версия: {camera['Аппаратная версия']}\n")
                        fw.write(f"Серийный номер: {camera['Серийный номер']}\n")
                        fw.write(f"MAC адрес: {camera['MAC адрес']}\n")
                        fw.write('-------------------------\n')
                
                messagebox.showinfo("Успех", f"Данные сохранены в {file_path}")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Ошибка при сохранении: {str(e)}")
    

    
    def _auto_adjust_columns(self, worksheet, df):
        """Автоматически подбирает ширину колонок"""
        try:
            from openpyxl.utils import get_column_letter
            
            for idx, column in enumerate(df.columns, 1):
                column_letter = get_column_letter(idx)
                
                # Определяем максимальную длину в колонке
                max_length = 0
                column_name = str(column)
                
                # Учитываем заголовок
                max_length = max(max_length, len(column_name))
                
                # Учитываем данные в ячейках
                for cell in worksheet[column_letter]:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        max_length = max(max_length, cell_length)
                
                # Добавляем небольшой запас
                adjusted_width = min(max_length + 2, 50)  # Максимальная ширина 50
                
                worksheet.column_dimensions[column_letter].width = adjusted_width
                
        except Exception as e:
            print(f"Ошибка авто-подбора колонок: {e}")

    def _apply_excel_formatting(self, worksheet, df):
        """Применяет форматирование к таблице"""
        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            # Стили
            header_font = Font(bold=True, color="FFFFFF", size=12)
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            centered = Alignment(horizontal="center", vertical="center")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Форматируем заголовки
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = centered
                cell.border = border
            
            # Форматируем данные
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
            
            # Центрируем IP и MAC адрес
            center_columns = ['IP', 'MAC адрес']
            for idx, column in enumerate(df.columns, 1):
                if column in center_columns:
                    column_letter = get_column_letter(idx)
                    for row in range(2, worksheet.max_row + 1):
                        cell = worksheet[f"{column_letter}{row}"]
                        cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Замораживаем заголовки
            worksheet.freeze_panes = "A2"
            
        except Exception as e:
            print(f"Ошибка применения форматирования: {e}")

    def save_to_excel(self):
        """Сохранение результатов в Excel с авто-форматированием"""
        if not self.data:
            messagebox.showwarning("Ошибка", "Нет данных для сохранения!")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="nits_scan_results.xlsx"
        )
        
        if file_path:
            try:
                df = pd.DataFrame(self.data)
                
                columns_order = [
                    'IP', 'Тип устройства', 'Модель', 
                    'Версия ПО', 'Аппаратная версия',
                    'Серийный номер', 'MAC адрес'
                ]
                
                # Оставляем только существующие колонки
                existing_columns = [col for col in columns_order if col in df.columns]
                df = df[existing_columns]
                
                # ════════════════════════════════════════════════════════════════════
                # АВТО-ФОРМАТИРОВАНИЕ 
                # ════════════════════════════════════════════════════════════════════
                with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                    df.to_excel(writer, index=False, sheet_name='Результаты сканирования')
                    
                    workbook = writer.book
                    worksheet = writer.sheets['Результаты сканирования']
                    
                    # Авто-подбор ширины колонок
                    self._auto_adjust_columns(worksheet, df)
                    
                    # Добавляем форматирование
                    self._apply_excel_formatting(worksheet, df)
                    
                    # Добавляем фильтры
                    worksheet.auto_filter.ref = worksheet.dimensions
                # ════════════════════════════════════════════════════════════════════
                
                messagebox.showinfo("Успех", f"Данные сохранены в {file_path}")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Ошибка при сохранении: {str(e)}")
        

if __name__ == "__main__":
    root = tk.Tk()
    app = NitsWindow(root)
    root.mainloop()
