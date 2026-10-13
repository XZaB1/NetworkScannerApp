import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext, filedialog
import threading
import json
import os
import ipaddress
import requests
from requests.auth import HTTPDigestAuth
from bs4 import BeautifulSoup
from utils.path_manager import PathManager
from ai.mac_analyzer import MacAnalyzer
from ai.mac_utils import normalize_mac
from ai.mac_model import MacClassifier
from datetime import datetime
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class AxisWindow:
    def __init__(self, master):
        self.master = master
        self.master.title("Сканирование камер Axis")
        self.master.geometry("1000x800")
        
        # Параметры по умолчанию
        self.username = ''
        self.password = ''
        self.stop_scan = False
        self.running = False
        self.ip_ranges = []
        self.selected_ranges = []
        self.data = []

        # Инициализация анализатора MAC
        self.mac_analyzer = MacAnalyzer(log_callback=self.add_to_ai_log)


        self.create_widgets()
        self.load_ranges()
        self.master.after(500, self.create_ai_log_window)
    
    def create_ai_log_window(self):
        """Создает отдельное окно для лога ИИ"""
        self.ai_log_window = tk.Toplevel(self.master)
        self.ai_log_window.title("Лог анализа MAC-адресов (ИИ)")
        self.ai_log_window.geometry("800x400")

        self.ai_log_window.protocol("WM_DELETE_WINDOW", lambda: None)  # Игнорировать закрытие
        self.ai_log_window.resizable(False, False)  
        
        self.ai_log = scrolledtext.ScrolledText(self.ai_log_window, wrap=tk.WORD)
        self.ai_log.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        ttk.Button(
            self.ai_log_window, 
            text="Очистить лог",
            command=lambda: self.ai_log.delete(1.0, tk.END)
        ).pack(pady=5)
    
    def add_to_ai_log(self, message):
        """Вывод в лог ИИ"""
        self.ai_log.insert(tk.END, f"{datetime.now().strftime('%H:%M:%S')} - {message}\n")
        self.ai_log.see(tk.END)
        self.ai_log.update_idletasks()

    def create_widgets(self):
        """Создание элементов интерфейса"""
        control_frame = ttk.Frame(self.master)
        control_frame.pack(pady=10, fill=tk.X)
        
        range_frame = ttk.LabelFrame(self.master, text="Управление диапазонами IP")
        range_frame.pack(pady=5, padx=10, fill=tk.BOTH, expand=True)
        
        self.canvas = tk.Canvas(range_frame)
        scrollbar = ttk.Scrollbar(range_frame, orient="vertical", command=self.canvas.yview)
        self.scrollable_frame = ttk.Frame(self.canvas)
        
        self.scrollable_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=scrollbar.set)
        
        self.canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # Элементы управления
        self.start_btn = ttk.Button(control_frame, text="Начать сканирование", command=self.start_scan_thread)
        self.start_btn.pack(side=tk.LEFT, padx=5)
        
        self.stop_btn = ttk.Button(control_frame, text="Остановить", command=self.stop_scan_process, state=tk.DISABLED)
        self.stop_btn.pack(side=tk.LEFT, padx=5)
        

        
        ttk.Button(control_frame, text="Добавить диапазон", command=self.add_range_dialog).pack(side=tk.RIGHT, padx=5)
        ttk.Button(control_frame, text="Удалить выбранное", command=self.remove_selected_ranges).pack(side=tk.RIGHT, padx=5)
        
        # Прогресс-бар
        self.progress = ttk.Progressbar(self.master, orient="horizontal", mode="determinate")
        self.progress.pack(fill=tk.X, padx=10, pady=5)
        
        # Лог и результаты
        notebook = ttk.Notebook(self.master)
        
        log_frame = ttk.Frame(notebook)
        self.log = scrolledtext.ScrolledText(log_frame, height=15, wrap=tk.WORD)
        self.log.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        notebook.add(log_frame, text="Лог выполнения")
        
        result_frame = ttk.Frame(notebook)
        self.result_text = scrolledtext.ScrolledText(result_frame, height=15, wrap=tk.WORD)
        self.result_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        notebook.add(result_frame, text="Результаты")
        
        notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        # Фрейм для кнопок сохранения
        save_frame = ttk.Frame(self.master)
        save_frame.pack(pady=10)
        
        ttk.Button(save_frame, text="Сохранить в TXT", command=self.save_to_txt).pack(side=tk.LEFT, padx=5)
        ttk.Button(save_frame, text="Сохранить в Excel", command=self.save_to_excel).pack(side=tk.LEFT, padx=5)
        ttk.Button(save_frame, text="Сохранить дубликаты MAC (AI)", command=self.save_ai_duplicates).pack(side=tk.LEFT, padx=5)
        
        self.update_range_checkboxes()
    
    def load_ranges(self):
        """Загрузка диапазонов IP из файла"""
        try:
            os.makedirs("config", exist_ok=True)
            if os.path.exists("config/ranges_axis.json"):
                with open("config/ranges_axis.json", "r") as f:
                    self.ip_ranges = json.load(f)
            else:
                self.ip_ranges = [
                    ["10.119.196.0", "10.119.196.253"],
                    ["10.119.197.0", "10.119.197.255"],
                    ["10.119.204.0", "10.119.204.255"],
                    ["10.119.216.0", "10.119.218.255"]
                ]
                self.save_ranges()
            
            self.selected_ranges = self.ip_ranges.copy()
            self.update_range_checkboxes()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось загрузить диапазоны: {str(e)}")
    
    def save_ranges(self):
        """Сохранение диапазонов в файл"""
        try:
            with open("config/ranges_axis.json", "w") as f:
                json.dump(self.ip_ranges, f, indent=4)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить диапазоны: {str(e)}")
    
    def update_range_checkboxes(self):
        """Обновление чекбоксов диапазонов"""
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        
        self.range_vars = []
        for i, (start, end) in enumerate(self.ip_ranges, 1):
            var = tk.BooleanVar(value=(start, end) in self.selected_ranges)
            self.range_vars.append((var, (start, end)))
            
            chk = ttk.Checkbutton(
                self.scrollable_frame,
                text=f"Диапазон {i}: {start} — {end}",
                variable=var,
                command=lambda v=var, r=(start, end): self.toggle_range(v, r)
            )
            chk.pack(anchor="w", padx=5, pady=2)
    
    def toggle_range(self, var, range):
        """Переключение состояния диапазона"""
        if var.get():
            if range not in self.selected_ranges:
                self.selected_ranges.append(range)
        else:
            if range in self.selected_ranges:
                self.selected_ranges.remove(range)
    
    def add_range_dialog(self):
        """Диалог добавления нового диапазона"""
        dialog = tk.Toplevel(self.master)
        dialog.title("Добавить диапазон IP")
        dialog.geometry("400x200")
        dialog.resizable(False, False)
        
        ttk.Label(dialog, text="Начальный IP:").pack(pady=5)
        start_entry = ttk.Entry(dialog)
        start_entry.pack(pady=5)
        
        ttk.Label(dialog, text="Конечный IP:").pack(pady=5)
        end_entry = ttk.Entry(dialog)
        end_entry.pack(pady=5)
        
        def save_range():
            start = start_entry.get().strip()
            end = end_entry.get().strip()
            
            try:
                ipaddress.ip_address(start)
                ipaddress.ip_address(end)
                
                new_range = [start, end]
                if new_range not in self.ip_ranges:
                    self.ip_ranges.append(new_range)
                    self.selected_ranges.append(new_range)
                    self.save_ranges()
                    self.update_range_checkboxes()
                    dialog.destroy()
                else:
                    messagebox.showwarning("Внимание", "Такой диапазон уже существует")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Некорректные IP-адреса: {str(e)}")
        
        ttk.Button(dialog, text="Добавить", command=save_range).pack(pady=10)
        start_entry.focus()
    
    def remove_selected_ranges(self):
        """Удаление выбранных диапазонов"""
        ranges_to_remove = []
        for var, rng in self.range_vars:
            if var.get():
                ranges_to_remove.append(list(rng))
        
        if not ranges_to_remove:
            messagebox.showwarning("Внимание", "Не выбраны диапазоны для удаления")
            return
        
        if not messagebox.askyesno("Подтверждение", f"Удалить {len(ranges_to_remove)} выбранных диапазонов?"):
            return
        
        self.ip_ranges = [r for r in self.ip_ranges if r not in ranges_to_remove]
        self.selected_ranges = [r for r in self.selected_ranges if r not in ranges_to_remove]
        self.save_ranges()
        self.update_range_checkboxes()
        messagebox.showinfo("Успех", f"Удалено {len(ranges_to_remove)} диапазонов")
    
    def start_scan_thread(self):
        """Запуск сканирования в отдельном потоке"""
        if not self.selected_ranges:
            messagebox.showwarning("Ошибка", "Выберите хотя бы один диапазон!")
            return
        
        if self.running:
            messagebox.showwarning("Внимание", "Сканирование уже выполняется!")
            return
        
        self.stop_scan = False
        self.running = True
        self.data = []
        self.log.delete("1.0", tk.END)
        self.result_text.delete("1.0", tk.END)
        
        self.start_btn.config(state=tk.DISABLED)
        self.stop_btn.config(state=tk.NORMAL)
        
        self.scan_thread = threading.Thread(target=self.run_scan, daemon=True)
        self.scan_thread.start()
    
    def stop_scan_process(self):
        """Остановка сканирования"""
        if not self.running:
            return
        
        self.stop_scan = True
        self.running = False
        self.log.insert(tk.END, "\n[INFO] Сканирование остановлено пользователем\n")
        self.start_btn.config(state=tk.NORMAL)
        self.stop_btn.config(state=tk.DISABLED)
    
    def run_scan(self):
        """Основная функция сканирования"""
        try:
            total_ips = sum(
                int(ipaddress.ip_address(end)) - int(ipaddress.ip_address(start)) + 1
                for start, end in self.selected_ranges
            )
            
            self.progress.config(maximum=total_ips)
            processed = 0
            
            for start, end in self.selected_ranges:
                start_ip = ipaddress.ip_address(start)
                end_ip = ipaddress.ip_address(end)
                
                for ip_int in range(int(start_ip), int(end_ip) + 1):
                    if self.stop_scan or not self.running:
                        self.stop_btn.config(state=tk.DISABLED)
                        self.start_btn.config(state=tk.NORMAL)
                        self.running = False
                        return
                    
                    ip_str = str(ipaddress.ip_address(ip_int))
                    self.log.insert(tk.END, f"[{processed+1}/{total_ips}] Проверка {ip_str}... ")
                    self.log.see(tk.END)
                    self.master.update()
                    
                    camera_info = self.get_camera_info(ip_str)
                    if camera_info:
                        self.data.append(camera_info)
                        self.result_text.insert(tk.END, f"Найдена камера: {ip_str}\n")
                        self.log.insert(tk.END, "OK - камера найдена\n")
                    else:
                        self.log.insert(tk.END, "Нет ответа\n")
                    
                    processed += 1
                    self.progress.config(value=processed)
                    self.master.update()
            
            self.start_btn.config(state=tk.NORMAL)
            self.stop_btn.config(state=tk.DISABLED)
            self.log.insert(tk.END, f"\n[INFO] Сканирование завершено! Найдено камер: {len(self.data)}\n")
            self.running = False
            
        except Exception as e:
            self.log.insert(tk.END, f"\n[ERROR] Критическая ошибка: {str(e)}\n")
            self.start_btn.config(state=tk.NORMAL)
            self.stop_btn.config(state=tk.DISABLED)
            self.running = False
    
    def get_camera_info(self, ip):
        """Получение информации о камере Axis с анализом MAC"""
        try:
            url = f'http://{ip}/axis-cgi/param.cgi?action=list'
            response = requests.get(url, auth=HTTPDigestAuth(self.username, self.password), timeout=5)
            response.raise_for_status()
            
            soup = BeautifulSoup(response.text, 'html.parser')
            camera_info = soup.get_text()
            lines = camera_info.splitlines()
            
            camera_data = {
                'IP': ip,
                'Версия': 'Не указано',
                'Имя устройства': 'Не указано',
                'Серийный номер': 'Не указано',
                'MAC адрес': 'Не указано',
                'Тип устройства': 'Axis',
                'MAC Анализ': {'is_duplicate': False}
            }

            for line in lines:
                if 'System.SerialNumber' in line:
                    camera_data['Серийный номер'] = line.split('=')[1].strip()
                elif 'MACAddress' in line:
                    mac = line.split('=')[1].strip()
                    camera_data['MAC адрес'] = mac
                    
                    if mac != 'Не указано':
                        try:
                            result = self.mac_analyzer.analyze(ip, mac, "Axis")
                            camera_data['MAC Анализ'] = result
                            
                            time_str = datetime.now().strftime('%H:%M:%S')
                            analysis_msg = [
                                f"[{time_str}] Анализ камеры {ip}",
                                f"MAC: {mac}",
                                f"Серийник: {camera_data['Серийный номер']}",
                                f"Прогноз: {'🚫 ДУБЛИКАТ' if result['is_duplicate'] else '🟢 Нормальный'}",
                                f"Уровень риска: {result.get('risk_score', 0.0):.2f}",
                            ]
                            
                            if result['is_duplicate']:
                                analysis_msg.append(f"Конфликты: {', '.join(result['conflict_ips'])}")
                            
                            self.add_to_ai_log("\n".join(analysis_msg) + "\n" + "━"*50 + "\n")
                            
                        except Exception as e:
                            error_msg = f"[{time_str}] Ошибка анализа MAC {mac}: {str(e)}"
                            self.add_to_ai_log(error_msg)
                            camera_data['MAC Анализ']['error'] = str(e)
                elif 'FirmwareVersion' in line:
                    camera_data['Версия'] = line.split('=')[1].strip()
                elif 'DeviceName' in line:
                    camera_data['Имя устройства'] = line.split('=')[1].strip()
            
            if camera_data['Серийный номер'] == 'Не указано':
                self.log.insert(tk.END, "Пропуск - серийный номер не найден\n")
                return None
                
            return camera_data
            
        except requests.HTTPError as http_err:
            if http_err.response.status_code == 401:
                self.log.insert(tk.END, f"[AUTH] Ошибка авторизации для {ip}\n")
            else:
                self.log.insert(tk.END, f"[HTTP] Ошибка {http_err.response.status_code} для {ip}\n")
        except Exception as e:
            self.log.insert(tk.END, f"[ERROR] Ошибка для {ip}: {str(e)}\n")
        
        return None
    
    def save_ai_duplicates(self):
        """Сохраняет Excel-отчет только с камерами, имеющими дубликаты MAC"""
        if not self.data:
            messagebox.showwarning("Ошибка", "Нет данных для анализа!")
            return
    
        duplicates = [
            cam for cam in self.data 
            if cam.get('MAC Анализ', {}).get('is_duplicate', False)
        ]
    
        if not duplicates:
            messagebox.showinfo("Результат", "Дубликаты MAC не обнаружены")
            return
    
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="axis_mac_duplicates.xlsx"
        )
    
        if not file_path:
            return
    
        try:
            report_data = []
            for cam in duplicates:
                analysis = cam['MAC Анализ']
                report_data.append({
                    'IP': cam['IP'],
                    'MAC адрес': cam['MAC адрес'],
                    'Уровень риска': f"{analysis['risk_score']:.2f}",
                    'Конфликтующие IP': ', '.join(analysis['conflict_ips']),
                    'Модель': cam['Имя устройства'],
                    'Серийный номер': cam['Серийный номер'],
                    'Версия прошивки': cam['Версия']
                })
            
            df = pd.DataFrame(report_data)
            column_order = [
                'IP', 'MAC адрес', 'Уровень риска', 'Конфликтующие IP',
                'Модель', 'Серийный номер', 'Версия прошивки'
            ]
            df = df[column_order]
            
            writer = pd.ExcelWriter(file_path, engine='xlsxwriter')
            df.to_excel(writer, index=False, sheet_name='Дубликаты MAC')
            
            workbook = writer.book
            worksheet = writer.sheets['Дубликаты MAC']
            
            red_format = workbook.add_format({'bg_color': '#FFC7CE'})
            yellow_format = workbook.add_format({'bg_color': '#FFEB9C'})
            
            for row_num in range(1, len(df)+1):
                risk = float(df.iloc[row_num-1]['Уровень риска'])
                if risk > 0.8:
                    worksheet.set_row(row_num, None, red_format)
                elif risk > 0.5:
                    worksheet.set_row(row_num, None, yellow_format)
            
            writer.close()
            
            messagebox.showinfo(
                "Успех", 
                f"Сохранено {len(duplicates)} дубликатов в:\n{file_path}"
            )
        except Exception as e:
            messagebox.showerror(
                "Ошибка", 
                f"Не удалось сохранить файл:\n{str(e)}"
            )
    
    def save_to_txt(self):
        """Сохранение результатов в TXT"""
        if not self.data:
            messagebox.showwarning("Ошибка", "Нет данных для сохранения!")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt")],
            initialfile="axis_scan_results.txt"
        )
        
        if file_path:
            try:
                with open(file_path, 'w', encoding='utf-8') as fw:
                    for camera in self.data:
                        fw.write(f"Информация о камере {camera['IP']}:\n")
                        fw.write(f"Версия: {camera['Версия']}\n")
                        fw.write(f"Имя устройства: {camera['Имя устройства']}\n")
                        fw.write(f"Серийный номер: {camera['Серийный номер']}\n")
                        fw.write(f"MAC адрес: {camera['MAC адрес']}\n")
                        fw.write('-------------------------\n')
                
                messagebox.showinfo("Успех", f"Данные сохранены в {file_path}")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Ошибка при сохранении: {str(e)}")

    def _auto_adjust_columns(self, worksheet, df):
        """Автоматически подбирает ширину колонок"""
        try:
            from openpyxl.utils import get_column_letter
            
            for idx, column in enumerate(df.columns, 1):
                column_letter = get_column_letter(idx)
                
                # Определяем максимальную длину в колонке
                max_length = 0
                column_name = str(column)
                
                # Учитываем заголовок
                max_length = max(max_length, len(column_name))
                
                # Учитываем данные в ячейках
                for cell in worksheet[column_letter]:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        max_length = max(max_length, cell_length)
                
                # Добавляем небольшой запас
                adjusted_width = min(max_length + 2, 50)  # Максимальная ширина 50
                
                worksheet.column_dimensions[column_letter].width = adjusted_width
                
        except Exception as e:
            print(f"Ошибка авто-подбора колонок: {e}")

    def _apply_excel_formatting(self, worksheet, df):
        """Применяет форматирование к таблице"""
        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            # Стили
            header_font = Font(bold=True, color="FFFFFF", size=12)
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            centered = Alignment(horizontal="center", vertical="center")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Форматируем заголовки
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = centered
                cell.border = border
            
            # Форматируем данные
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
            
            # Центрируем IP и MAC адрес
            center_columns = ['IP', 'MAC адрес']
            for idx, column in enumerate(df.columns, 1):
                if column in center_columns:
                    column_letter = get_column_letter(idx)
                    for row in range(2, worksheet.max_row + 1):
                        cell = worksheet[f"{column_letter}{row}"]
                        cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Замораживаем заголовки
            worksheet.freeze_panes = "A2"
            
        except Exception as e:
            print(f"Ошибка применения форматирования: {e}")
    
    def save_to_excel(self):
        """Сохранение результатов в Excel с авто-форматированием"""
        if not self.data:
            messagebox.showwarning("Ошибка", "Нет данных для сохранения!")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="axis_scan_results.xlsx"
        )
        
        if file_path:
            try:
                df = pd.DataFrame(self.data)
                
                # Упорядочиваем колонки
                columns_order = ['IP', 'Тип устройства', 'Версия', 
                                'Имя устройства', 'Серийный номер', 'MAC адрес']
                
                # Оставляем только существующие колонки
                existing_columns = [col for col in columns_order if col in df.columns]
                df = df[existing_columns]
                
                # ════════════════════════════════════════════════════════════════════
                # АВТО-ФОРМАТИРОВАНИЕ КАК В ДРУГИХ СКАНЕРАХ
                # ════════════════════════════════════════════════════════════════════
                with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                    df.to_excel(writer, index=False, sheet_name='Результаты сканирования')
                    
                    workbook = writer.book
                    worksheet = writer.sheets['Результаты сканирования']
                    
                    # Авто-подбор ширины колонок
                    self._auto_adjust_columns(worksheet, df)
                    
                    # Добавляем форматирование
                    self._apply_excel_formatting(worksheet, df)
                    
                    # Добавляем фильтры
                    worksheet.auto_filter.ref = worksheet.dimensions
                # ════════════════════════════════════════════════════════════════════
                
                messagebox.showinfo("Успех", f"Данные сохранены в {file_path}")
            except Exception as e:
                messagebox.showerror("Ошибка", f"Ошибка при сохранении: {str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    app = AxisWindow(root)
    root.mainloop()
